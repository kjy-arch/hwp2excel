"""hwp2excel 단위 테스트 — 한글 설치나 샘플 파일 없이 실행된다.

    python -m unittest
"""

import os
import tempfile
import threading
import time
import unittest
import zipfile
from unittest import mock

import openpyxl

import hwp2excel
from hwp2excel import (
    App,
    convert_cell_value,
    extract_tables_via_hwpx,
    filter_tables_by_page,
    make_merge,
    make_sheet_name,
    pick_table_title,
    tables_to_excel,
)


NS = 'xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph"'


def cell_xml(row, col, row_span, col_span, text):
    return (
        f'<hp:tc><hp:subList><hp:p><hp:run><hp:t>{text}</hp:t></hp:run>'
        f'</hp:p></hp:subList>'
        f'<hp:cellAddr colAddr="{col}" rowAddr="{row}"/>'
        f'<hp:cellSpan colSpan="{col_span}" rowSpan="{row_span}"/></hp:tc>'
    )


def make_hwpx(section_xml):
    """섹션 XML 한 개를 담은 HWPX 파일을 만들고 경로를 돌려준다."""
    fd, path = tempfile.mkstemp(suffix=".hwpx")
    os.close(fd)
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("Contents/section0.xml", section_xml)
    return path


# 제목 문단 + 자료출처 문단 + 표. 표는 실제 HWPX처럼 문단 안에 들어간다.
# 제목 2행 + 자료 1행. 구분(A1:A2 세로 병합), 대상자(B1:C1 가로 병합).
SAMPLE_SECTION = (
    f'<hp:sec {NS}>'
    '<hp:p><hp:run><hp:t>1-1-3. 병역준비역 등 자원 현황</hp:t></hp:run></hp:p>'
    '<hp:p><hp:run><hp:t>자료출처：병역판정검사과</hp:t></hp:run></hp:p>'
    '<hp:p><hp:run><hp:tbl rowCnt="3" colCnt="3">'
    '<hp:tr>'
    + cell_xml(0, 0, 2, 1, "구분")
    + cell_xml(0, 1, 1, 2, "대상자")
    + '</hp:tr><hp:tr>'
    + cell_xml(1, 1, 1, 1, "소계")
    + cell_xml(1, 2, 1, 1, "19세자")
    + '</hp:tr><hp:tr>'
    + cell_xml(2, 0, 1, 1, "계")
    + cell_xml(2, 1, 1, 1, "923,444")
    + cell_xml(2, 2, 1, 1, "256,517")
    + '</hp:tr></hp:tbl></hp:run></hp:p></hp:sec>'
)


class MakeMergeTest(unittest.TestCase):
    def test_single_cell_is_not_a_merge(self):
        self.assertIsNone(make_merge(0, 0, 1, 1, 5, 5))

    def test_column_merge(self):
        self.assertEqual(make_merge(0, 1, 1, 3, 5, 5), (0, 1, 0, 3))

    def test_row_merge(self):
        self.assertEqual(make_merge(0, 0, 2, 1, 5, 5), (0, 0, 1, 0))

    def test_clamped_to_table_bounds(self):
        self.assertEqual(make_merge(1, 1, 9, 9, 3, 3), (1, 1, 2, 2))

    def test_zero_span_treated_as_one(self):
        self.assertIsNone(make_merge(2, 2, 0, 0, 5, 5))


class ExtractHwpxTest(unittest.TestCase):
    def setUp(self):
        self.path = make_hwpx(SAMPLE_SECTION)
        self.addCleanup(os.remove, self.path)

    def test_title_comes_from_the_paragraph_above_the_table(self):
        (tbl,) = extract_tables_via_hwpx(self.path)
        self.assertEqual(tbl["title"], "1-1-3. 병역준비역 등 자원 현황")

    def test_grid_and_merges(self):
        (tbl,) = extract_tables_via_hwpx(self.path)
        self.assertEqual(tbl["row_count"], 3)
        self.assertEqual(tbl["col_count"], 3)
        self.assertEqual(tbl["rows"][0], ["구분", "대상자", ""])
        self.assertEqual(tbl["rows"][1], ["", "소계", "19세자"])
        self.assertEqual(sorted(tbl["merges"]), [(0, 0, 1, 0), (0, 1, 0, 2)])

    def test_rejects_non_hwpx(self):
        fd, path = tempfile.mkstemp(suffix=".hwpx")
        os.close(fd)
        self.addCleanup(os.remove, path)
        with zipfile.ZipFile(path, "w") as zf:
            zf.writestr("mimetype", "application/hwp+zip")
        with self.assertRaises(ValueError):
            extract_tables_via_hwpx(path)


class TablesToExcelTest(unittest.TestCase):
    def write(self, tables):
        fd, path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        self.addCleanup(os.remove, path)
        tables_to_excel(tables, path)
        return path

    def sheet(self, tables, index=0):
        return openpyxl.load_workbook(self.write(tables)).worksheets[index]

    def test_merges_are_applied(self):
        path = make_hwpx(SAMPLE_SECTION)
        self.addCleanup(os.remove, path)
        ws = self.sheet(extract_tables_via_hwpx(path))
        self.assertEqual(
            sorted(str(r) for r in ws.merged_cells.ranges),
            ["A1:A2", "B1:C1"],
        )
        self.assertEqual(ws["A1"].value, "구분")
        self.assertEqual(ws["B1"].value, "대상자")
        self.assertEqual(ws["B2"].value, "소계")
        self.assertEqual(ws["B3"].value, 923444)

    def test_header_block_covers_vertically_merged_first_row(self):
        path = make_hwpx(SAMPLE_SECTION)
        self.addCleanup(os.remove, path)
        ws = self.sheet(extract_tables_via_hwpx(path))
        # 첫 행이 2행까지 병합됐으므로 2행까지가 제목
        self.assertTrue(ws["B2"].font.bold)
        self.assertEqual(ws["B2"].alignment.horizontal, "center")
        self.assertFalse(ws["B3"].font.bold)

    def test_single_header_row_without_merges(self):
        tables = [{
            "index": 1, "row_count": 2, "col_count": 2,
            "rows": [["가", "나"], ["1", "2"]], "merges": [],
        }]
        ws = self.sheet(tables)
        self.assertEqual(len(ws.merged_cells.ranges), 0)
        self.assertTrue(ws["A1"].font.bold)
        self.assertFalse(ws["A2"].font.bold)

    def test_layout_table_merged_top_to_bottom_keeps_one_header_row(self):
        # 첫 열이 표 전체 높이로 병합된 도형용 표 — 전체가 제목이 되면 안 된다
        tables = [{
            "index": 1, "row_count": 3, "col_count": 2,
            "rows": [["가", "1"], ["", "2"], ["", "3"]],
            "merges": [(0, 0, 2, 0)],
        }]
        ws = self.sheet(tables)
        self.assertTrue(ws["B1"].font.bold)
        self.assertFalse(ws["B2"].font.bold)

    def test_missing_merges_key_is_tolerated(self):
        tables = [{
            "index": 1, "row_count": 1, "col_count": 2,
            "rows": [["가", "나"]],
        }]
        ws = self.sheet(tables)
        self.assertEqual(ws["B1"].value, "나")

    def test_overlapping_merges_are_skipped(self):
        tables = [{
            "index": 1, "row_count": 2, "col_count": 3,
            "rows": [["가", "", ""], ["1", "2", "3"]],
            "merges": [(0, 0, 0, 2), (0, 1, 1, 2)],
        }]
        ws = self.sheet(tables)
        self.assertEqual([str(r) for r in ws.merged_cells.ranges], ["A1:C1"])

    def test_text_starting_with_equals_is_not_a_formula(self):
        tables = [{
            "index": 1, "row_count": 1, "col_count": 2,
            "rows": [["=1+1", "= 합계"]], "merges": [],
        }]
        ws = self.sheet(tables)
        self.assertEqual(ws["A1"].value, "=1+1")
        self.assertEqual(ws["A1"].data_type, "s")
        self.assertEqual(ws["B1"].data_type, "s")

    def test_wide_header_does_not_stretch_its_column(self):
        long_text = "아" * 60
        tables = [{
            "index": 1, "row_count": 2, "col_count": 3,
            "rows": [["", long_text, ""], ["1", "2", "3"]],
            "merges": [(0, 1, 0, 2)],
        }]
        ws = self.sheet(tables)
        self.assertLess(ws.column_dimensions["B"].width, 10)


class PickTitleTest(unittest.TestCase):
    def test_numbered_heading_wins_over_nearer_source_note(self):
        self.assertEqual(
            pick_table_title(["앞 문단", "1-1-3. 병역준비역 등 자원 현황",
                              "자료출처：병역판정검사과"]),
            "1-1-3. 병역준비역 등 자원 현황",
        )

    def test_notes_are_skipped(self):
        for noise in ["자료출처 : 기획재정담당관", "※ 주의", "(단위 : 명)",
                      "주 : 반올림", "* 비고"]:
            self.assertEqual(pick_table_title(["통 계 설 명 자 료", noise]),
                             "통 계 설 명 자 료")

    def test_nearest_line_when_nothing_is_numbered(self):
        self.assertEqual(pick_table_title(["먼 문단", "가까운 문단"]),
                         "가까운 문단")

    def test_none_when_there_is_nothing_usable(self):
        self.assertIsNone(pick_table_title([]))
        self.assertIsNone(pick_table_title(["자료출처：어딘가", ""]))


class SheetNameTest(unittest.TestCase):
    def name(self, used=None, **tbl):
        tbl.setdefault("index", 1)
        return make_sheet_name(tbl, used if used is not None else set())

    def test_page_and_title(self):
        self.assertEqual(
            self.name(index=9, page=25, title="1-1-3. 병역준비역 등 자원 현황"),
            "25_1-1-3. 병역준비역 등 자원 현황")

    def test_table_number_stands_in_when_page_is_unknown(self):
        self.assertEqual(self.name(index=9, title="자원 현황"),
                         "표009_자원 현황")

    def test_falls_back_to_table_number_without_a_title(self):
        self.assertEqual(self.name(index=1, page=1), "표001")
        self.assertEqual(self.name(index=7, title=None), "표007")

    def test_truncated_to_31_characters(self):
        name = self.name(index=10, page=26, title="가" * 60)
        self.assertEqual(len(name), 31)
        self.assertTrue(name.startswith("26_"))
        self.assertTrue(name.endswith("…"))

    def test_forbidden_characters_are_removed(self):
        name = self.name(index=2, page=3, title="가/나:다*라?마[바]사\\아")
        self.assertEqual(name, "3_가 나 다 라 마 바 사 아")

    def test_duplicate_names_get_a_number(self):
        used = set()
        first = self.name(used, index=5, page=23, title="병역자원 관리 현황")
        second = self.name(used, index=6, page=23, title="병역자원 관리 현황")
        self.assertEqual(first, "23_병역자원 관리 현황")
        self.assertEqual(second, "23_병역자원 관리 현황_2")

    def test_duplicate_of_a_truncated_name_still_fits(self):
        used = set()
        long = "1-1-4. 병역준비역 편입대상자 조사 현황 - 청별"
        a = self.name(used, index=1, page=26, title=long)
        b = self.name(used, index=2, page=26, title=long)
        self.assertLessEqual(len(a), 31)
        self.assertLessEqual(len(b), 31)
        self.assertNotEqual(a, b)


class FilterByPageTest(unittest.TestCase):
    def setUp(self):
        self.tables = [
            {"index": 1, "page": 1},
            {"index": 9, "page": 25},
            {"index": 10, "page": 26},
            {"index": 11, "page": 27},
        ]

    def test_range_is_inclusive_on_both_ends(self):
        picked = filter_tables_by_page(self.tables, 25, 27)
        self.assertEqual([t["index"] for t in picked], [9, 10, 11])

    def test_single_page(self):
        picked = filter_tables_by_page(self.tables, 26, 26)
        self.assertEqual([t["index"] for t in picked], [10])

    def test_original_table_numbers_are_kept(self):
        picked = filter_tables_by_page(self.tables, 25, 25)
        self.assertEqual(picked[0]["index"], 9)

    def test_empty_when_no_table_on_those_pages(self):
        self.assertEqual(filter_tables_by_page(self.tables, 2, 16), [])

    def test_missing_page_info_raises(self):
        with self.assertRaises(ValueError):
            filter_tables_by_page([{"index": 1}], 1, 5)


class PageInputUiTest(unittest.TestCase):
    def setUp(self):
        try:
            self.app = App()
        except Exception as e:  # 디스플레이 없는 환경
            self.skipTest(f"Tk를 띄울 수 없음: {e}")
        self.app.withdraw()
        self.addCleanup(self.app.destroy)

    def test_default_is_whole_document_with_inputs_off(self):
        self.assertEqual(self.app.page_mode_var.get(), "all")
        self.assertEqual(str(self.app.first_entry["state"]), "disabled")
        self.assertIsNone(self.app._page_range())

    def test_range_mode_enables_inputs(self):
        self.app.page_mode_var.set("range")
        self.app._sync_page_inputs()
        self.assertEqual(str(self.app.first_entry["state"]), "normal")
        self.app.first_var.set("25")
        self.app.last_var.set("27")
        self.assertEqual(self.app._page_range(), (25, 27))

    def test_direct_parsing_disables_the_whole_row(self):
        self.app.page_mode_var.set("range")
        self.app.method_var.set("ole")
        self.app._sync_page_inputs()
        self.assertEqual(str(self.app.rb_all["state"]), "disabled")
        self.assertEqual(str(self.app.first_entry["state"]), "disabled")
        self.assertEqual(self.app.page_hint["text"], "한글 연동 전용")
        self.assertIsNone(self.app._page_range())

    def test_bad_page_input_is_rejected(self):
        self.app.page_mode_var.set("range")
        for first, last in [("가", "3"), ("", ""), ("5", "2"), ("0", "3")]:
            self.app.first_var.set(first)
            self.app.last_var.set(last)
            with self.assertRaises(ValueError):
                self.app._page_range()


class ConvertThreadingTest(unittest.TestCase):
    """변환은 작업 스레드에서, Tk 호출은 모두 메인 스레드에서."""

    def setUp(self):
        try:
            self.app = App()
        except Exception as e:  # 디스플레이 없는 환경
            self.skipTest(f"Tk를 띄울 수 없음: {e}")
        self.app.withdraw()
        self.addCleanup(self.app.destroy)

        fd, self.src = tempfile.mkstemp(suffix=".hwp")
        os.close(fd)
        self.addCleanup(os.remove, self.src)
        self.app.hwp_var.set(self.src)
        self.app.out_var.set(self.src + ".xlsx")
        self.app.method_var.set("ole")

    def run_conversion(self, tables):
        main = threading.current_thread()
        seen = {}

        def extract(path):
            seen["extract"] = threading.current_thread()
            if isinstance(tables, Exception):
                raise tables
            return tables

        def to_excel(tbls, out, progress_cb=None):
            for i in range(len(tbls)):
                progress_cb(i + 1, len(tbls))

        def dialog(kind):
            def show(*args, **kwargs):
                seen[kind] = threading.current_thread()
            return show

        with mock.patch.object(hwp2excel, "extract_tables_via_olefile",
                               extract), \
                mock.patch.object(hwp2excel, "tables_to_excel", to_excel), \
                mock.patch.object(hwp2excel.messagebox, "showinfo",
                                  dialog("info")), \
                mock.patch.object(hwp2excel.messagebox, "showwarning",
                                  dialog("warning")), \
                mock.patch.object(hwp2excel.messagebox, "showerror",
                                  dialog("error")):
            self.app._start()
            deadline = time.time() + 5
            while str(self.app.btn["state"]) == "disabled":
                self.assertLess(time.time(), deadline, "변환이 끝나지 않음")
                self.app.update()
                time.sleep(0.01)
        return main, seen

    def test_extraction_runs_off_the_main_thread(self):
        main, seen = self.run_conversion([{"index": 1}])
        self.assertIsNot(seen["extract"], main)

    def test_dialogs_are_shown_on_the_main_thread(self):
        main, seen = self.run_conversion([{"index": 1}, {"index": 2}])
        self.assertIs(seen["info"], main)
        self.assertEqual(self.app.status_var.get(), "완료! 표 2개 변환됨")
        self.assertEqual(self.app.progress["value"], 0)

    def test_warning_when_no_tables(self):
        main, seen = self.run_conversion([])
        self.assertIs(seen["warning"], main)

    def test_error_is_reported_on_the_main_thread(self):
        main, seen = self.run_conversion(RuntimeError("깨진 파일"))
        self.assertIs(seen["error"], main)
        self.assertEqual(self.app.status_var.get(), "오류 발생")


class ConvertCellValueTest(unittest.TestCase):
    def test_thousands_separator_becomes_int(self):
        self.assertEqual(convert_cell_value("923,444"), (923444, None))

    def test_percent_becomes_ratio_with_format(self):
        value, num_format = convert_cell_value("18.6%")
        self.assertAlmostEqual(value, 0.186)
        self.assertEqual(num_format, "0.0%")

    def test_text_with_unit_stays_text(self):
        self.assertEqual(convert_cell_value("6,666명"), ("6,666명", None))


if __name__ == "__main__":
    unittest.main()
